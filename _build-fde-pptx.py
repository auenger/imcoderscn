import glob
import os
from pptx import Presentation
from pptx.util import Inches

D = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(D, "fde-slides-screenshots")
OUT = os.path.join(D, "fde-million-salary-slides-V1.1.pptx")

prs = Presentation()
# 16:9 宽屏，与截图比例一致
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]  # 空白版式

imgs = sorted(glob.glob(os.path.join(IMG_DIR, "slide-*.png")))
assert imgs, "no screenshots found"

for img in imgs:
    slide = prs.slides.add_slide(blank)
    # 每页一张图，铺满整页（左上角对齐，宽高=幻灯片尺寸）
    slide.shapes.add_picture(img, 0, 0, prs.slide_width, prs.slide_height)

prs.save(OUT)
print("saved:", OUT)
print("slides:", len(imgs))
print("size : 13.333 x 7.5 in (16:9)")
